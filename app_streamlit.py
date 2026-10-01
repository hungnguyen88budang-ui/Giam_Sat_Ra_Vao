import glob
import json
import os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st

# --- CẤU HÌNH THƯ MỤC ---
DATA_DIR = r"D:\GiamSat_RaVao\data_nhanvien"
EXCEL_BAOCAO_PATH = r"D:\GiamSat_RaVao\GiamSat_RaVao.xlsx"
MAPPING_JSON = r"D:\GiamSat_RaVao\danh_sach_ten.json"

os.makedirs(DATA_DIR, exist_ok=True)

st.set_page_config(
    page_title="Hệ Thống Giám Sát Ra Vào", page_icon="🛡️", layout="wide"
)

st.title("🛡️ HỆ THỐNG GIÁM SÁT RA VÀO CÔNG NHÂN")

tab1, tab2, tab3 = st.tabs([
    "📋 Danh Sách & Xóa Nhân Sự",
    "➕ Thêm Công Nhân Mới",
    "📊 Nhật Ký & Báo Cáo Vi Phạm",
])


def load_name_mapping():
  if os.path.exists(MAPPING_JSON):
    try:
      with open(MAPPING_JSON, "r", encoding="utf-8") as f:
        return json.load(f)
    except:
      pass
  return {}


def save_name_mapping(mapping):
  with open(MAPPING_JSON, "w", encoding="utf-8") as f:
    json.dump(mapping, f, ensure_ascii=False, indent=4)


def load_data_from_images():
  main_files = glob.glob(os.path.join(DATA_DIR, "*_main.*"))
  mapping = load_name_mapping()
  data = []

  for file_path in main_files:
    filename = os.path.basename(file_path)
    ma_cn = filename.split("_main.")[0]
    ho_ten = mapping.get(ma_cn, ma_cn)

    ext = filename.split(".")[-1]
    path_left = os.path.join(DATA_DIR, f"{ma_cn}_left.{ext}")
    path_right = os.path.join(DATA_DIR, f"{ma_cn}_right.{ext}")

    data.append({
        "Mã Công Nhân": ma_cn,
        "Họ và Tên": ho_ten,
        "Ảnh Chính Diện": file_path,
        "Ảnh Góc Trái": (
            file_path if not os.path.exists(path_left) else path_left
        ),
        "Ảnh Góc Phải": (
            file_path if not os.path.exists(path_right) else path_right
        ),
        "Trạng Thái": "Đã ghi nhận AI",
    })

  return pd.DataFrame(data)


def delete_worker(ma_cn):
  pattern = os.path.join(DATA_DIR, f"{ma_cn}_*")
  for f_path in glob.glob(pattern):
    try:
      os.remove(f_path)
    except Exception:
      pass

  mapping = load_name_mapping()
  if ma_cn in mapping:
    del mapping[ma_cn]
    save_name_mapping(mapping)


def style_excel_file(filename):
  try:
    wb = openpyxl.load_workbook(filename)
    ws = wb.active
    ws.freeze_panes = "A2"

    header_fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )

    for col in range(1, ws.max_column + 1):
      cell = ws.cell(row=1, column=col)
      cell.fill = header_fill
      cell.font = header_font
      cell.alignment = center_align

    for row in range(2, ws.max_row + 1):
      for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = left_align if col in [2, 8] else center_align

    for col in ws.columns:
      max_len = 0
      col_letter = openpyxl.utils.get_column_letter(col[0].column)
      for cell in col:
        val_str = str(cell.value or "")
        if len(val_str) > max_len:
          max_len = len(val_str)
      ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(filename)
  except Exception:
    pass


# TAB 1: DANH SÁCH & XÓA
with tab1:
  col_head, col_btn = st.columns([4, 1])
  with col_head:
    st.subheader("📋 Danh Sách Nhân Sự Quản Lý")
  with col_btn:
    if st.button("🔄 Tải lại dữ liệu"):
      st.rerun()

  df_nv = load_data_from_images()

  if df_nv.empty:
    st.info("📂 Thư mục `data_nhanvien` hiện chưa có dữ liệu công nhân!")
  else:
    st.dataframe(
        df_nv,
        column_config={
            "Ảnh Chính Diện": st.column_config.ImageColumn(
                "Ảnh Chính Diện", width="small"
            ),
            "Ảnh Góc Trái": st.column_config.ImageColumn(
                "Ảnh Góc Trái", width="small"
            ),
            "Ảnh Góc Phải": st.column_config.ImageColumn(
                "Ảnh Góc Phải", width="small"
            ),
        },
        hide_index=True,
        width="stretch",
    )

    st.write("---")
    st.subheader("🗑️ Xóa Công Nhân Khỏi Hệ Thống")
    list_ma_cn = df_nv["Mã Công Nhân"].tolist()
    col_sel, col_del = st.columns([3, 1])

    with col_sel:
      selected_ma = st.selectbox(
          "Chọn Mã Công Nhân cần xóa:",
          options=list_ma_cn,
          format_func=lambda x: f"{x} - {load_name_mapping().get(x, x)}",
      )

    with col_del:
      st.write(" ")
      st.write(" ")
      if st.button("❌ Xóa Công Nhân", type="primary"):
        if selected_ma:
          delete_worker(selected_ma)
          st.success(
              f"✅ Đã xóa hoàn toàn dữ liệu và ảnh của mã **{selected_ma}**!"
          )
          st.rerun()

# TAB 2: THÊM MỚI
with tab2:
  st.subheader("➕ Tải Ảnh & Đăng Ký Họ Tên Công Nhân")

  with st.form("form_them_nv", clear_on_submit=True):
    col_a, col_b = st.columns(2)
    with col_a:
      ma_cn = st.text_input("1. Mã Công Nhân (Ví dụ: CN001):").strip()
    with col_b:
      ho_ten = st.text_input("2. Họ và Tên Công Nhân:").strip()

    st.write("---")
    st.write("📸 **Tải lên bộ ảnh nhận diện**")
    c1, c2, c3 = st.columns(3)
    with c1:
      img_main = st.file_uploader(
          "Ảnh chính diện (Bắt buộc)", type=["jpg", "png", "jpeg"]
      )
    with c2:
      img_left = st.file_uploader("Ảnh góc trái", type=["jpg", "png", "jpeg"])
    with c3:
      img_right = st.file_uploader("Ảnh góc Phải", type=["jpg", "png", "jpeg"])

    submit = st.form_submit_button("💾 Lưu Dữ Liệu Công Nhân")

    if submit:
      if not ma_cn or not ho_ten:
        st.error("⚠️ Vui lòng nhập đầy đủ Mã Công Nhân và Họ Tên!")
      elif not img_main:
        st.error("⚠️ Bắt buộc phải chọn Ảnh chính diện!")
      else:
        mapping = load_name_mapping()
        mapping[ma_cn] = ho_ten
        save_name_mapping(mapping)

        with open(os.path.join(DATA_DIR, f"{ma_cn}_main.jpg"), "wb") as f:
          f.write(img_main.getbuffer())
        if img_left:
          with open(os.path.join(DATA_DIR, f"{ma_cn}_left.jpg"), "wb") as f:
            f.write(img_left.getbuffer())
        if img_right:
          with open(os.path.join(DATA_DIR, f"{ma_cn}_right.jpg"), "wb") as f:
            f.write(img_right.getbuffer())

        st.success(f"✅ Đã thêm thành công: **{ho_ten}** ({ma_cn})!")
        st.rerun()

# TAB 3: NHẬT KÝ EXCEL & LỌC XÓA NGÀY CŨ
with tab3:
  st.subheader("📊 Nhật Ký Giám Sát & Báo Cáo Vi Phạm")

  if not os.path.exists(EXCEL_BAOCAO_PATH):
    st.warning("⚠️ Chưa có dữ liệu lịch sử ra/vào!")
  else:
    try:
      df_log = pd.read_excel(EXCEL_BAOCAO_PATH)

      if df_log.empty:
        st.info("📂 File nhật ký hiện đang trống.")
      else:
        # Trích xuất danh sách Ngày từ cột Giờ RA / Giờ VÀO
        df_log["Ngay_Filter"] = (
            df_log["Giờ RA"]
            .fillna(df_log["Giờ VÀO"])
            .astype(str)
            .str.split(" ")
            .str[0]
        )
        available_days = sorted(
            [d for d in df_log["Ngay_Filter"].unique() if len(d) == 10],
            reverse=True,
        )

        # THANH LỌC DỮ LIỆU BẰNG CỘT TÌM KIẾM & BỘ LỌC NGÀY
        c_date, c_search, c_dl = st.columns([2, 3, 2])

        with c_date:
          selected_date = st.selectbox(
              "📅 Lọc Theo Ngày:", options=["Tất cả"] + available_days
          )

        with c_search:
          search_term = st.text_input(
              "🔍 Tìm Tên / Mã công nhân:"
          ).strip()

        with c_dl:
          st.write(" ")
          with open(EXCEL_BAOCAO_PATH, "rb") as f:
            st.download_button(
                label="📥 TẢI FILE EXCEL (.XLSX)",
                data=f,
                file_name="BaoCao_RaVao_CongNhan.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        # Lọc dataframe hiển thị
        df_display = df_log.copy()
        if selected_date != "Tất cả":
          df_display = df_display[df_display["Ngay_Filter"] == selected_date]

        if search_term:
          df_display = df_display[
              df_display["Tên Nhân Viên"]
              .astype(str)
              .str.contains(search_term, case=False)
          ]

        # Hiển thị bảng dữ liệu
        st.dataframe(
            df_display.drop(columns=["Ngay_Filter"]),
            width="stretch",
            hide_index=True,
        )

        st.write("---")

        # KHU VỰC XÓA DỮ LIỆU
        st.subheader("🗑️ Quản Lý Xóa Lượt Ghi Nhận")
        col_del_row, col_del_day = st.columns(2)

        # Xóa 1 dòng lượt cụ thể
        with col_del_row:
          st.write("**Xóa 1 lượt ra/vào cụ thể:**")
          list_ma_luot = df_display["Mã Lượt"].tolist()
          if list_ma_luot:
            selected_luot = st.selectbox(
                "Chọn Mã Lượt cần xóa:", options=list_ma_luot
            )
            if st.button("❌ Xóa Lượt Này", type="primary"):
              df_log_new = df_log[df_log["Mã Lượt"] != selected_luot].drop(
                  columns=["Ngay_Filter"]
              )
              df_log_new.to_excel(EXCEL_BAOCAO_PATH, index=False)
              style_excel_file(EXCEL_BAOCAO_PATH)
              st.success(f"✅ Đã xóa thành công Mã Lượt: {selected_luot}!")
              st.rerun()

        # Xóa dữ liệu cả ngày cũ
        with col_del_day:
          st.write("**Xóa nhanh dữ liệu theo ngày:**")
          if selected_date != "Tất cả":
            if st.button(
                f"🔥 Xóa Tất Cả Dữ Liệu Ngày {selected_date}", type="primary"
            ):
              df_log_new = df_log[
                  df_log["Ngay_Filter"] != selected_date
              ].drop(columns=["Ngay_Filter"])
              df_log_new.to_excel(EXCEL_BAOCAO_PATH, index=False)
              style_excel_file(EXCEL_BAOCAO_PATH)
              st.success(f"✅ Đã xóa toàn bộ dữ liệu ngày {selected_date}!")
              st.rerun()
          else:
            st.info("💡 Hãy chọn 1 ngày cụ thể ở ô **Lọc Theo Ngày** phía trên để mở tính năng xóa cả ngày.")

    except Exception as e:
      st.error(f"Lỗi đọc/ghi Excel: {e}")